import logging
from pathlib import Path
from typing import List
from pptx import Presentation
from .base import BaseExtractor, ExtractedChunk

logger = logging.getLogger(__name__)

class PPTXExtractor(BaseExtractor):
    def extract(self, file_path: Path) -> List[ExtractedChunk]:
        chunks: List[ExtractedChunk] = []
        try:
            prs = Presentation(str(file_path))
            for slide_idx, slide in enumerate(prs.slides, start=1):
                slide_texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            p_text = paragraph.text.strip()
                            if p_text:
                                slide_texts.append(p_text)
                
                full_slide_text = "\n".join(slide_texts).strip()
                if full_slide_text:
                    chunks.append(
                        ExtractedChunk(
                            text=full_slide_text,
                            page_number=None,
                            slide_number=slide_idx,
                            snippet=full_slide_text[:250],
                            chunk_index=len(chunks)
                        )
                    )
        except Exception as e:
            logger.error(f"Erro ao extrair texto do PPTX {file_path}: {e}")
        return chunks
